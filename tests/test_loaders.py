from dataclasses import replace
from pathlib import Path
from unittest.mock import MagicMock, patch
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document

from agentic_rag.ingestion_limits import DEFAULT_INGESTION_LIMITS, IngestionLimitError
from agentic_rag.loaders import UnsupportedDocumentError, load_document


def test_loads_text_file(tmp_path: Path) -> None:
    path = tmp_path / "notes.txt"
    path.write_text("  Agentic retrieval uses tools.  ", encoding="utf-8")

    sections = load_document(path)

    assert [section.text for section in sections] == ["Agentic retrieval uses tools."]
    assert sections[0].source == "notes.txt"


def test_loads_docx_file(tmp_path: Path) -> None:
    path = tmp_path / "guide.docx"
    document = Document()
    document.add_paragraph("First paragraph.")
    document.add_paragraph("Second paragraph.")
    document.save(str(path))

    sections = load_document(path)

    assert sections[0].text == "First paragraph.\n\nSecond paragraph."
    assert sections[0].source == "guide.docx"


def test_loads_non_empty_pdf_pages(tmp_path: Path) -> None:
    path = tmp_path / "paper.pdf"
    path.touch()
    first_page = MagicMock()
    first_page.extract_text.return_value = "Page one"
    empty_page = MagicMock()
    empty_page.extract_text.return_value = None

    with patch("agentic_rag.loaders.PdfReader") as reader:
        reader.return_value.pages = [first_page, empty_page]
        sections = load_document(path)

    assert len(sections) == 1
    assert sections[0].text == "Page one"
    assert sections[0].page == 1


def test_rejects_source_larger_than_upload_limit(tmp_path: Path) -> None:
    path = tmp_path / "large.txt"
    path.write_bytes(b"12345")
    limits = replace(DEFAULT_INGESTION_LIMITS, max_source_bytes=4)

    with pytest.raises(IngestionLimitError, match="upload limit"):
        load_document(path, limits=limits)


def test_rejects_extracted_text_over_character_limit(tmp_path: Path) -> None:
    path = tmp_path / "large.txt"
    path.write_text("123456", encoding="utf-8")
    limits = replace(DEFAULT_INGESTION_LIMITS, max_extracted_characters=5)

    with pytest.raises(IngestionLimitError, match="Document text"):
        load_document(path, limits=limits)


def test_rejects_docx_with_oversized_uncompressed_xml(tmp_path: Path) -> None:
    path = tmp_path / "large.docx"
    with ZipFile(path, mode="w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "x" * 100)
    limits = replace(
        DEFAULT_INGESTION_LIMITS,
        max_docx_document_xml_bytes=50,
    )

    with pytest.raises(IngestionLimitError, match="document text expands"):
        load_document(path, limits=limits)


def test_rejects_pdf_with_too_many_pages(tmp_path: Path) -> None:
    path = tmp_path / "large.pdf"
    path.touch()
    limits = replace(DEFAULT_INGESTION_LIMITS, max_sections=1)

    with (
        patch("agentic_rag.loaders.PdfReader") as reader,
        pytest.raises(IngestionLimitError, match="more than 1 pages"),
    ):
        reader.return_value.pages = [MagicMock(), MagicMock()]
        load_document(path, limits=limits)


def test_rejects_pdf_with_excessive_decompressed_content(tmp_path: Path) -> None:
    path = tmp_path / "large.pdf"
    path.touch()
    page = MagicMock()
    page.get_contents.return_value.get_data.return_value = b"123456"
    limits = replace(
        DEFAULT_INGESTION_LIMITS,
        max_pdf_stream_bytes=5,
        max_pdf_decompressed_bytes=5,
    )

    with (
        patch("agentic_rag.loaders.PdfReader") as reader,
        pytest.raises(IngestionLimitError, match="PDF content expands"),
    ):
        reader.return_value.pages = [page]
        load_document(path, limits=limits)


def test_rejects_unsupported_document(tmp_path: Path) -> None:
    path = tmp_path / "archive.zip"
    path.touch()

    with pytest.raises(UnsupportedDocumentError, match="\\.zip"):
        load_document(path)
